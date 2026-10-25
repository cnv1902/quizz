import csv
from io import BytesIO, StringIO
from typing import List, Optional
from sqlalchemy.orm import Session
import openpyxl

from app.crud import crud_measure_word
from app.models.measure_word import MeasureWord


def get_all(db: Session) -> List[MeasureWord]:
    return crud_measure_word.get_all(db)


def get_by_id(db: Session, item_id: int) -> Optional[MeasureWord]:
    return crud_measure_word.get_by_id(db, item_id)


def create(db: Session, vietnamese: str, pinyin: str, hanzi: str, order_index: int = 0) -> MeasureWord:
    return crud_measure_word.create(db, vietnamese, pinyin, hanzi, order_index)


def bulk_create(db: Session, items: List[dict]) -> List[MeasureWord]:
    return crud_measure_word.bulk_create(db, items)


def update(db: Session, item_id: int, **kwargs) -> Optional[MeasureWord]:
    return crud_measure_word.update(db, item_id, **kwargs)


def delete(db: Session, item_id: int) -> bool:
    return crud_measure_word.delete(db, item_id)


def bulk_delete(db: Session, ids: List[int]) -> int:
    return crud_measure_word.bulk_delete(db, ids)


def generate_sample_excel_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lượng từ mẫu"
    ws.append(["Tiếng Việt", "Pinyin", "Chữ Hán"])
    sample_rows = [
        ["Lượng từ chung dùng cho người và vật", "gè", "个"],
        ["Lượng từ dùng cho sách, vở dạng quyển", "běn", "本"],
        ["Lượng từ dùng cho đồ vật có cán/tay cầm (ô, ghế, dao)", "bǎ", "把"],
    ]
    for row in sample_rows:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def import_from_file(db: Session, filename: str, file_bytes: bytes) -> List[MeasureWord]:
    lower_name = (filename or "").lower()
    raw_cards = []

    if lower_name.endswith((".xlsx", ".xls")):
        wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("File Excel không có dữ liệu")
        start_row = 1 if any(str(c).strip().lower() in ("tiếng việt", "vietnamese", "nghĩa") for c in rows[0] if c) else 0
        for row in rows[start_row:]:
            if not row or all(v is None for v in row):
                continue
            v = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
            p = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            h = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            if v and h:
                raw_cards.append({"vietnamese": v, "pinyin": p, "hanzi": h})
    else:
        content = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.reader(StringIO(content))
        rows = list(reader)
        if not rows:
            raise ValueError("File CSV không có dữ liệu")
        start_row = 1 if any(c.strip().lower() in ("tiếng việt", "vietnamese", "nghĩa") for c in rows[0] if c) else 0
        for row in rows[start_row:]:
            if not row or all(not v.strip() for v in row):
                continue
            v = row[0].strip() if len(row) > 0 else ""
            p = row[1].strip() if len(row) > 1 else ""
            h = row[2].strip() if len(row) > 2 else ""
            if v and h:
                raw_cards.append({"vietnamese": v, "pinyin": p, "hanzi": h})

    if not raw_cards:
        raise ValueError("Không tìm thấy dòng dữ liệu hợp lệ (cần có ít nhất cột Tiếng Việt và Chữ Hán)")

    created = crud_measure_word.bulk_create(db, raw_cards)
    return created
