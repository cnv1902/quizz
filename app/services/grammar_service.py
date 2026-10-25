import csv
from io import BytesIO, StringIO
from typing import List, Optional
from sqlalchemy.orm import Session
import openpyxl

from app.crud import crud_grammar_point
from app.models.grammar_point import GrammarPoint


def get_all(db: Session) -> List[GrammarPoint]:
    return crud_grammar_point.get_all(db)


def get_by_id(db: Session, item_id: int) -> Optional[GrammarPoint]:
    return crud_grammar_point.get_by_id(db, item_id)


def create(db: Session, vietnamese: str, pinyin: str, hanzi: str, order_index: int = 0) -> GrammarPoint:
    return crud_grammar_point.create(db, vietnamese, pinyin, hanzi, order_index)


def bulk_create(db: Session, items: List[dict]) -> List[GrammarPoint]:
    return crud_grammar_point.bulk_create(db, items)


def update(db: Session, item_id: int, **kwargs) -> Optional[GrammarPoint]:
    return crud_grammar_point.update(db, item_id, **kwargs)


def delete(db: Session, item_id: int) -> bool:
    return crud_grammar_point.delete(db, item_id)


def bulk_delete(db: Session, ids: List[int]) -> int:
    return crud_grammar_point.bulk_delete(db, ids)


def generate_sample_excel_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ngữ pháp mẫu"
    ws.append(["Giải thích Công thức", "Pinyin", "Chữ Hán"])
    sample_rows = [
        ["Cấu trúc hỏi phương thức: Chủ ngữ + 怎么 + Động từ?", "zěnme", "怎么"],
        ["Trợ từ kết cấu biểu thị quan hệ sở hữu: Thành phần bổ nghĩa + 的 + Danh từ", "de", "的"],
    ]
    for row in sample_rows:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def import_from_file(db: Session, filename: str, file_bytes: bytes) -> List[GrammarPoint]:
    lower_name = (filename or "").lower()
    raw_cards = []
    header_keywords = (
        "giải thích công thức", "giai thich cong thuc",
        "giải thích", "giai thich",
        "công thức", "cong thuc",
        "tiếng việt", "tieng viet", "vietnamese", "nghĩa", "nghia",
        "chữ hán", "chu han", "hanzi", "pinyin"
    )

    if lower_name.endswith((".xlsx", ".xls")):
        wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("File Excel không có dữ liệu")

        col_v_idx, col_p_idx, col_h_idx = 0, 1, 2
        first_row_str = " ".join(str(c or "").lower() for c in rows[0])
        has_header = any(k in first_row_str for k in header_keywords)

        if has_header:
            for idx, cell in enumerate(rows[0]):
                c_str = str(cell or "").strip().lower()
                if any(k in c_str for k in ("giải thích", "giai thich", "công thức", "cong thuc", "tiếng việt", "tieng viet", "nghĩa", "nghia")):
                    col_v_idx = idx
                elif any(k in c_str for k in ("pinyin", "phiên âm", "phien am")):
                    col_p_idx = idx
                elif any(k in c_str for k in ("chữ hán", "chu han", "hán tự", "hanzi", "cấu trúc", "cau truc")):
                    col_h_idx = idx
            data_rows = rows[1:]
        else:
            data_rows = rows

        for row in data_rows:
            if not row or all(v is None for v in row):
                continue
            v = str(row[col_v_idx]).strip() if len(row) > col_v_idx and row[col_v_idx] is not None else ""
            p = str(row[col_p_idx]).strip() if len(row) > col_p_idx and row[col_p_idx] is not None else ""
            h = str(row[col_h_idx]).strip() if len(row) > col_h_idx and row[col_h_idx] is not None else ""
            if v and h:
                raw_cards.append({"vietnamese": v, "pinyin": p, "hanzi": h})
    else:
        content = file_bytes.decode("utf-8-sig", errors="replace")
        delimiter = ","
        first_line = content.splitlines()[0] if content.splitlines() else ""
        if "\t" in first_line:
            delimiter = "\t"
        elif ";" in first_line and first_line.count(";") >= first_line.count(","):
            delimiter = ";"
        reader = csv.reader(StringIO(content), delimiter=delimiter)
        rows = list(reader)
        if not rows:
            raise ValueError("File CSV không có dữ liệu")

        col_v_idx, col_p_idx, col_h_idx = 0, 1, 2
        first_row_str = " ".join(str(c or "").lower() for c in rows[0])
        has_header = any(k in first_row_str for k in header_keywords)

        if has_header:
            for idx, cell in enumerate(rows[0]):
                c_str = str(cell or "").strip().lower()
                if any(k in c_str for k in ("giải thích", "giai thich", "công thức", "cong thuc", "tiếng việt", "tieng viet", "nghĩa", "nghia")):
                    col_v_idx = idx
                elif any(k in c_str for k in ("pinyin", "phiên âm", "phien am")):
                    col_p_idx = idx
                elif any(k in c_str for k in ("chữ hán", "chu han", "hán tự", "hanzi", "cấu trúc", "cau truc")):
                    col_h_idx = idx
            data_rows = rows[1:]
        else:
            data_rows = rows

        for row in data_rows:
            if not row or all(not str(v).strip() for v in row):
                continue
            v = str(row[col_v_idx]).strip() if len(row) > col_v_idx and row[col_v_idx] is not None else ""
            p = str(row[col_p_idx]).strip() if len(row) > col_p_idx and row[col_p_idx] is not None else ""
            h = str(row[col_h_idx]).strip() if len(row) > col_h_idx and row[col_h_idx] is not None else ""
            if v and h:
                raw_cards.append({"vietnamese": v, "pinyin": p, "hanzi": h})

    if not raw_cards:
        raise ValueError("Không tìm thấy dòng dữ liệu hợp lệ (cần có ít nhất cột Giải thích Công thức và Chữ Hán)")

    created = crud_grammar_point.bulk_create(db, raw_cards)
    return created
