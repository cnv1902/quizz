import sys
import os
import re
from pathlib import Path
import openpyxl

# Thiết lập encoding UTF-8 cho console
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

try:
    import jieba
    import pypinyin
except ImportError:
    print("Vui lòng cài đặt jieba và pypinyin: pip install jieba pypinyin")
    sys.exit(1)

# Bản đồ chuẩn hóa dấu câu
PUNCTUATION_MAP = {
    "！": "!",
    "？": "?",
    "。": ".",
    "，": ",",
    "：": ":",
    "；": ";",
    "、": ",",
    "“": '"',
    "”": '"',
    "‘": "'",
    "’": "'",
    "（": "(",
    "）": ")",
}

# Các từ đặc biệt cần thanh nhẹ hoặc phiên âm chuẩn theo ngữ cảnh HSK
SPECIAL_WORDS_PINYIN = {
    "什么": "shénme",
    "名字": "míngzi",
    "学生": "xuésheng",
    "老师": "lǎoshī",
    "认识": "rènshi",
    "喜欢": "xǐhuan",
    "漂亮": "piàoliang",
    "朋友": "péngyou",
    "衣服": "yīfu",
    "东西": "dōngxi",
    "爸爸": "bàba",
    "妈妈": "māma",
    "哥哥": "gēge",
    "弟弟": "dìdi",
    "姐姐": "jiějie",
    "妹妹": "mèimei",
    "儿子": "érzi",
    "女儿": "nǚ'ér",
    "他们": "tāmen",
    "她们": "tāmen",
    "我们": "wǒmen",
    "你们": "nǐmen",
    "这里": "zhèlǐ",
    "那里": "nàlǐ",
    "哪里": "nǎlǐ",
}


def convert_chinese_sentence_to_pinyin(sentence: str) -> str:
    """
    Chuyển đổi một câu tiếng Trung thành Pinyin chuẩn tự nhiên:
    - Phân tách cụm từ bằng jieba.
    - Ghép pinyin cho từng từ ghép (ví dụ: 老师 -> lǎoshī, 什么 -> shénme).
    - Chuẩn hóa dấu câu gắn liền vào từ đứng trước.
    """
    if not sentence:
        return ""

    clean_text = str(sentence).strip()
    words = jieba.cut(clean_text)
    parts = []

    for w in words:
        w_str = str(w).strip()
        if not w_str:
            continue

        # Xử lý dấu câu
        if w_str in PUNCTUATION_MAP:
            punc = PUNCTUATION_MAP[w_str]
            if parts:
                parts[-1] = parts[-1] + punc
            else:
                parts.append(punc)
        elif w_str in ["!", "?", ".", ",", ":", ";", "(", ")", '"', "'"]:
            if parts:
                parts[-1] = parts[-1] + w_str
            else:
                parts.append(w_str)
        else:
            # Kiểm tra từ điển đặc biệt
            if w_str in SPECIAL_WORDS_PINYIN:
                parts.append(SPECIAL_WORDS_PINYIN[w_str])
            else:
                # Dùng pypinyin sinh pinyin có thanh điệu
                py_list = pypinyin.lazy_pinyin(w_str, style=pypinyin.Style.TONE)
                parts.append("".join(py_list))

    # Ghép lại thành câu pinyin hoàn chỉnh
    result = " ".join(parts)
    return result


def process_excel_file(file_path: str, output_path: str = None) -> bool:
    """
    Xử lý file Excel: tìm cột Tiếng Trung, sinh Pinyin và chèn vào cột Pinyin.
    """
    path = Path(file_path)
    if not path.exists():
        print(f"Lỗi: Không tìm thấy file tại '{file_path}'")
        return False

    out_file = output_path or file_path
    print(f"--> Đang mở file Excel: {path.resolve()}")

    wb = openpyxl.load_workbook(path)
    ws = wb.active

    # Xác định vị trí các cột từ dòng 1 (header)
    col_chinese = None
    col_pinyin = None
    col_viet = None

    for col in range(1, ws.max_column + 1):
        val = str(ws.cell(1, col).value or "").strip().lower()
        if "trung" in val or "chinese" in val or "hán" in val:
            col_chinese = col
        elif "pinyin" in val or "phiên âm" in val:
            col_pinyin = col
        elif "việt" in val or "vietnamese" in val or "nghĩa" in val:
            col_viet = col

    # Nếu không tìm thấy qua tên cột, tự động phát hiện theo thứ tự thông dụng (3 cột: Cột 1 tiếng Trung, Cột 2 Pinyin, Cột 3 Tiếng Việt)
    if not col_chinese:
        col_chinese = 1
    if not col_pinyin:
        # Nếu cột 2 có header hoặc rỗng
        col_pinyin = 2
    if not col_viet:
        col_viet = 3

    print(f"    Cột Tiếng Trung: Cột {col_chinese} ('{ws.cell(1, col_chinese).value}')")
    print(f"    Cột Pinyin     : Cột {col_pinyin} ('{ws.cell(1, col_pinyin).value}')")
    print(f"    Cột Tiếng Việt : Cột {col_viet} ('{ws.cell(1, col_viet).value}')")

    # Đảm bảo header cột Pinyin có tên
    if not ws.cell(1, col_pinyin).value:
        ws.cell(1, col_pinyin, "Pinyin")

    count = 0
    start_row = 2
    for r in range(start_row, ws.max_row + 1):
        chinese_text = ws.cell(r, col_chinese).value
        if chinese_text:
            pinyin_text = convert_chinese_sentence_to_pinyin(str(chinese_text))
            ws.cell(r, col_pinyin, pinyin_text)
            count += 1

    wb.save(out_file)
    print(f"\n✅ Đã điền Pinyin thành công cho {count} câu!")
    print(f"📁 Đã lưu file: {Path(out_file).resolve()}")

    # In thử 5 dòng đầu
    print("\n--- Mẫu 5 dòng sau khi chèn Pinyin ---")
    wb_check = openpyxl.load_workbook(out_file, data_only=True)
    ws_check = wb_check.active
    for r in range(1, min(7, ws_check.max_row + 1)):
        ch = ws_check.cell(r, col_chinese).value
        py = ws_check.cell(r, col_pinyin).value
        vi = ws_check.cell(r, col_viet).value
        print(f"Dòng {r}: [{ch}] | [{py}] | [{vi}]")

    return True


if __name__ == "__main__":
    target_file = r"d:\quizz\data\hanziquizz_ghep_cau_import_hsk1.xlsx"
    if len(sys.argv) > 1:
        target_file = sys.argv[1]

    output_file = None
    if len(sys.argv) > 2:
        output_file = sys.argv[2]

    process_excel_file(target_file, output_file)
